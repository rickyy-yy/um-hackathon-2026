"""Report export — PDF (reportlab), XLSX (openpyxl), DOCX (python-docx)."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from app.models.report import Report

PRIMARY = "#0F6E56"
ACCENT = "#F0DD62"
ALERT = "#D64933"
SURFACE = "#C6DABF"
BG = "#F3E9D2"


def export_pdf(report: Report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "title", parent=styles["Title"], textColor=colors.HexColor(PRIMARY)
    )
    heading_style = ParagraphStyle(
        "h", parent=styles["Heading2"], textColor=colors.HexColor(PRIMARY)
    )
    body_style = styles["BodyText"]

    flow: list[Any] = []
    flow.append(Paragraph(report.title, title_style))
    flow.append(Spacer(1, 0.3 * cm))
    flow.append(
        Paragraph(
            f"Generated: {report.generated_at.strftime('%d %b %Y %H:%M')}",
            body_style,
        )
    )
    flow.append(Spacer(1, 0.5 * cm))

    s = report.summary_json
    agg = s.get("aggregate_metrics", {})
    flow.append(Paragraph("Ringkasan", heading_style))
    flow.append(
        Paragraph(
            f"Total Revenue: RM{agg.get('total_revenue', 0):,.2f}<br/>"
            f"Total COGS: RM{agg.get('total_cogs', 0):,.2f}<br/>"
            f"Gross Profit: RM{agg.get('gross_profit', 0):,.2f}<br/>"
            f"Margin: {agg.get('overall_margin_pct', 0):.1f}%<br/>"
            f"Items Sold: {agg.get('total_items_sold', 0)}",
            body_style,
        )
    )
    flow.append(Spacer(1, 0.5 * cm))

    flow.append(Paragraph("Menu Item Breakdown", heading_style))
    data = [["Item", "Units", "Revenue", "Cost", "Profit", "Margin %", "Score"]]
    for row in s.get("menu_item_breakdown", []):
        data.append(
            [
                row["item_name"],
                row["units_sold"],
                f"RM{row['total_revenue']:,.2f}",
                f"RM{row['total_cost']:,.2f}",
                f"RM{row['gross_profit']:,.2f}",
                f"{row['margin_pct']:.1f}%",
                row["profitability_score"].title(),
            ]
        )
    table = Table(data, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(PRIMARY)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cccccc")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    flow.append(table)
    flow.append(Spacer(1, 0.5 * cm))

    flow.append(Paragraph("Cadangan AI (Kira)", heading_style))
    for para in (report.ai_recommendations or "").split("\n\n"):
        if para.strip():
            flow.append(Paragraph(para.replace("\n", "<br/>"), body_style))
            flow.append(Spacer(1, 0.2 * cm))

    tax = s.get("tax_estimation")
    if tax:
        flow.append(Paragraph("Anggaran Cukai LHDN", heading_style))
        flow.append(
            Paragraph(
                f"Estimated taxable income: RM{tax['estimated_taxable_income']:,.2f}<br/>"
                f"Estimated tax: RM{tax['estimated_tax']:,.2f}<br/>"
                f"Bracket: {tax['tax_bracket']}<br/>"
                f"<i>{tax['note']}</i>",
                body_style,
            )
        )

    doc.build(flow)
    return buf.getvalue()


def export_xlsx(report: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"

    hdr = Font(color="FFFFFF", bold=True)
    hdr_fill = PatternFill(start_color=PRIMARY.lstrip("#"), end_color=PRIMARY.lstrip("#"), fill_type="solid")

    s = report.summary_json
    agg = s.get("aggregate_metrics", {})

    summary_rows = [
        ("Metric", "Value"),
        ("Report", report.title),
        ("Total Revenue (RM)", agg.get("total_revenue", 0)),
        ("Total COGS (RM)", agg.get("total_cogs", 0)),
        ("Gross Profit (RM)", agg.get("gross_profit", 0)),
        ("Margin (%)", agg.get("overall_margin_pct", 0)),
        ("Items Sold", agg.get("total_items_sold", 0)),
    ]
    for row in summary_rows:
        summary_ws.append(row)
    for cell in summary_ws[1]:
        cell.font = hdr
        cell.fill = hdr_fill

    items_ws = wb.create_sheet("Menu Items")
    items_ws.append(
        ["Item", "Category", "Units", "Revenue", "Cost", "Profit", "Margin %", "Score"]
    )
    for cell in items_ws[1]:
        cell.font = hdr
        cell.fill = hdr_fill
    for row in s.get("menu_item_breakdown", []):
        items_ws.append(
            [
                row["item_name"],
                row.get("category") or "",
                row["units_sold"],
                row["total_revenue"],
                row["total_cost"],
                row["gross_profit"],
                row["margin_pct"],
                row["profitability_score"],
            ]
        )

    tax_ws = wb.create_sheet("Tax Estimation")
    tax = s.get("tax_estimation", {})
    for k, v in tax.items():
        tax_ws.append([k, v])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_docx(report: Report) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    title = doc.add_heading(report.title, level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x0F, 0x6E, 0x56)

    s = report.summary_json
    agg = s.get("aggregate_metrics", {})

    doc.add_heading("Ringkasan", level=1)
    para = doc.add_paragraph()
    para.add_run(
        f"Total Revenue: RM{agg.get('total_revenue', 0):,.2f}\n"
        f"Total COGS: RM{agg.get('total_cogs', 0):,.2f}\n"
        f"Gross Profit: RM{agg.get('gross_profit', 0):,.2f}\n"
        f"Margin: {agg.get('overall_margin_pct', 0):.1f}%\n"
        f"Items Sold: {agg.get('total_items_sold', 0)}"
    )

    doc.add_heading("Menu Item Breakdown", level=1)
    table = doc.add_table(rows=1, cols=7)
    hdr = table.rows[0].cells
    for i, head in enumerate(["Item", "Units", "Revenue", "Cost", "Profit", "Margin %", "Score"]):
        hdr[i].text = head
    for row in s.get("menu_item_breakdown", []):
        cells = table.add_row().cells
        cells[0].text = row["item_name"]
        cells[1].text = str(row["units_sold"])
        cells[2].text = f"RM{row['total_revenue']:,.2f}"
        cells[3].text = f"RM{row['total_cost']:,.2f}"
        cells[4].text = f"RM{row['gross_profit']:,.2f}"
        cells[5].text = f"{row['margin_pct']:.1f}%"
        cells[6].text = row["profitability_score"]

    doc.add_heading("Cadangan AI (Kira)", level=1)
    for para_text in (report.ai_recommendations or "").split("\n\n"):
        if para_text.strip():
            doc.add_paragraph(para_text)

    tax = s.get("tax_estimation")
    if tax:
        doc.add_heading("Anggaran Cukai LHDN", level=1)
        doc.add_paragraph(
            f"Estimated taxable income: RM{tax['estimated_taxable_income']:,.2f}\n"
            f"Estimated tax: RM{tax['estimated_tax']:,.2f}\n"
            f"Bracket: {tax['tax_bracket']}\n"
            f"{tax['note']}"
        )

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
